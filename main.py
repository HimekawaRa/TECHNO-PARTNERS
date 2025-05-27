import shutil
import tempfile
from fastapi import FastAPI, File, UploadFile, HTTPException, Form
import os
import re
import subprocess
import json
import logging
import openai
from dotenv import load_dotenv
from docx import Document
from docx.document import Document as _Document
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph
from pdf2image import convert_from_path
import base64
from fastapi.staticfiles import StaticFiles

from app.auto_parser import split_questions_logic, pipeline_mcq, \
    build_rows_with_placeholders, clean_math_and_sub, pipeline_matching

# Настройка логирования
logging.basicConfig(level=logging.INFO, format='%(asctime)s %(levelname)s %(message)s')
logger = logging.getLogger(__name__)

load_dotenv()
openai.api_key = os.getenv("OPENAI_API_KEY")

app = FastAPI()



BASE_DIR = os.path.dirname(__file__)
STATIC_DIR = os.path.join(BASE_DIR, "static")
IMG_DIR = os.path.join(STATIC_DIR, "img")
os.makedirs(IMG_DIR, exist_ok=True)
app.mount(
    "/img",
    StaticFiles(directory=os.path.join(STATIC_DIR, "img")),
    name="img",
)




@app.get("/healthcheck")
async def healthcheck():
    return {"status": "ok"}

@app.post("/split-multiple-choice-questions/", tags=["Python Parser"])
async def split_questions_api(
    file: UploadFile = File(...),
    subject_name: str = Form(..., description="Название предмета"),
    subject_namekz: str = Form(..., description="Название предмета на казахском"),
    language: str = Form("рус", description="Язык задания"),
    klass: str = Form(..., description="Класс, например '10 ЕМН'"),
    tip: int = Form(1, description="Тип задания (целое число)")
):
    tmp = tempfile.mkdtemp()
    filename = (file.filename or "input.docx").replace(' ', '_')
    src = os.path.join(tmp, filename)
    docname = os.path.splitext(filename)[0]
    try:
        # 1) Сохраняем загруженный .docx во временную папку
        with open(src, "wb") as f:
            f.write(await file.read())

        # 2) Разбираем документ на вопросы и извлекаем медиа
        try:
            raw_list = split_questions_logic(src)

            # Перемещаем извлечённые изображения в static/img/<docname>
            media_dir = os.path.join(tmp, "media")
            target_img_dir = os.path.join(IMG_DIR, docname)
            os.makedirs(target_img_dir, exist_ok=True)

            if os.path.isdir(media_dir):
                for root, _, files in os.walk(media_dir):
                    for fname in files:
                        src_img = os.path.join(root, fname)
                        fname_clean = fname.replace(' ', '_')
                        dst_img = os.path.join(target_img_dir, fname_clean)
                        shutil.copy2(src_img, dst_img)

        except Exception as e:
            logger.error("Ошибка при split_questions_logic: %s", e)
            raise HTTPException(status_code=400, detail=str(e))

        # 3) Прогоним через пайплайн для multiple choice
        states = [pipeline_mcq(raw_item) for raw_item in raw_list]

        # 4) Собираем итоговые строки
        subject = {"name": subject_name, "namekz": subject_namekz}
        db_rows = build_rows_with_placeholders(states, subject, language, klass, tip)

        # 5) Чистим LaTeX/математические выражения
        for row in db_rows:
            row["vopros"] = clean_math_and_sub(row.get("vopros", ""))
            row["exp"] = clean_math_and_sub(row.get("exp", ""))
            row["otvety"] = [clean_math_and_sub(opt) for opt in row.get("otvety", [])]

        return {"questions": db_rows}

    finally:
        shutil.rmtree(tmp, ignore_errors=True)

@app.post("/split-matching-questions/", tags=["Python Parser"])
async def split_questions_api(
    file: UploadFile = File(...),
    subject_name: str = Form(..., description="Название предмета"),
    subject_namekz: str = Form(..., description="Название предмета на казахском"),
    language: str = Form("рус", description="Язык задания"),
    klass: str = Form(..., description="Класс, например '10 ЕМН'"),
    tip: int = Form(1, description="Тип задания (целое число)")
):
    tmp = tempfile.mkdtemp()
    filename = (file.filename or "input.docx").replace(' ', '_')
    src = os.path.join(tmp, filename)
    docname = os.path.splitext(filename)[0]
    try:
        # 1) Сохраняем загруженный .docx во временную папку
        with open(src, "wb") as f:
            f.write(await file.read())

        # 2) Разбираем документ на вопросы и извлекаем медиа
        try:
            raw_list = split_questions_logic(src)

            # Перемещаем все извлечённые Pandoc'ом медиа-файлы в static/img/<docname>
            media_dir = os.path.join(tmp, "media")
            target_img_dir = os.path.join(IMG_DIR, docname)
            os.makedirs(target_img_dir, exist_ok=True)

            if os.path.isdir(media_dir):
                for root, _, files in os.walk(media_dir):
                    for fname in files:
                        src_img = os.path.join(root, fname)
                        fname_clean = fname.replace(' ', '_')
                        dst_img = os.path.join(target_img_dir, fname_clean)
                        shutil.copy2(src_img, dst_img)

        except Exception as e:
            logger.error("Ошибка при split_questions_logic: %s", e)
            raise HTTPException(status_code=400, detail=str(e))

        # 3) Прогоним через нашу ML-пайплайн логику
        states = [pipeline_matching(raw_item) for raw_item in raw_list]

        # 4) Собираем итоговые строки
        subject = {"name": subject_name, "namekz": subject_namekz}
        db_rows = build_rows_with_placeholders(states, subject, language, klass, tip)

        # 5) Чистим математические выражения
        for row in db_rows:
            row["vopros"] = clean_math_and_sub(row.get("vopros", ""))
            row["exp"]    = clean_math_and_sub(row.get("exp", ""))
            for group_name, opts in row.get("otvety", {}).items():
                for key, val in list(opts.items()):
                    opts[key] = clean_math_and_sub(val)

        return {"questions": db_rows}

    finally:
        shutil.rmtree(tmp, ignore_errors=True)